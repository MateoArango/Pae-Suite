from __future__ import annotations

import random
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path("fixtures/student-data-alright - repeated - attendant  - x4.xlsx")
OUTPUT = Path(
    "fixtures/student-data-alright - repeated - attendant  - x100 - correct importation.xlsx"
)
FIRST_DATA_ROW = 3
RECORD_COUNT = 100
LAST_DATA_ROW = FIRST_DATA_ROW + RECORD_COUNT - 1
SEED = 20260803
UNSUPPORTED_DOCUMENT_TYPES = {"NIUP", "NES", "PPT", "RUMV", "PEP", "PA"}
UNSUPPORTED_GRADES = {-1, 0, 15}

FIRST_NAMES = [
    "Adriana", "Alejandro", "Andrés", "Camila", "Carlos", "Daniel", "Daniela",
    "David", "Elena", "Emilia", "Esteban", "Felipe", "Gabriela", "Isabella",
    "Javier", "Juliana", "Laura", "Lucía", "Manuel", "Mariana", "Mateo",
    "Natalia", "Nicolás", "Paula", "Samuel", "Sara", "Santiago", "Sofía",
    "Valentina", "Tomás",
]

SECOND_NAMES = [
    "Alejandra", "Antonio", "Cristina", "Eduardo", "Fernanda", "Francisco",
    "Isabel", "José", "Luis", "Marcela", "María", "Miguel", "Patricia",
    "Rafael", "Victoria",
]

LAST_NAMES = [
    "Álvarez", "Cárdenas", "Castillo", "Díaz", "Flórez", "García", "Gómez",
    "González", "Gutiérrez", "Hernández", "Jiménez", "López", "Martínez",
    "Mendoza", "Moreno", "Muñoz", "Navarro", "Ortiz", "Pérez", "Ramírez",
    "Rojas", "Romero", "Sánchez", "Suárez", "Torres", "Vargas",
]


def validation_values(sheet, cell_coordinate: str) -> list[str]:
    for validation in sheet.data_validations.dataValidation:
        if cell_coordinate in validation.cells:
            formula = validation.formula1
            if isinstance(formula, str) and formula.startswith('"') and formula.endswith('"'):
                return formula[1:-1].split(",")
    raise RuntimeError(f"No canonical list validation found for {cell_coordinate}")


def email_part(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    ascii_value = normalized.encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "", ascii_value.lower())


def unique_numbers(rng: random.Random, start: int, end: int, count: int) -> list[int]:
    return rng.sample(range(start, end + 1), count)


def main() -> None:
    rng = random.Random(SEED)
    workbook = load_workbook(SOURCE)
    sheet = workbook[workbook.sheetnames[0]]

    canonical_document_types = validation_values(sheet, "A3")
    canonical_grades = [int(value) for value in validation_values(sheet, "G3")]
    canonical_population_types = validation_values(sheet, "I3")

    beneficiary_document_types = [
        value
        for value in canonical_document_types
        if value not in UNSUPPORTED_DOCUMENT_TYPES
    ]
    supported_grades = [value for value in canonical_grades if value not in UNSUPPORTED_GRADES]

    # Correct the dropdown rules so users cannot select values rejected by import.
    for validation in sheet.data_validations.dataValidation:
        if "A3" in validation.cells:
            validation.formula1 = f'"{",".join(beneficiary_document_types)}"'
        if "G3" in validation.cells:
            validation.formula1 = f'"{",".join(str(value) for value in supported_grades)}"'

    beneficiary_documents = unique_numbers(rng, 10_000_000, 99_999_999, RECORD_COUNT)
    beneficiary_phones = unique_numbers(rng, 300_000_0000, 399_999_9999, RECORD_COUNT)

    attendant_first_name = rng.choice(FIRST_NAMES)
    attendant_last_name = rng.choice(LAST_NAMES)
    attendant_document = rng.randint(100_000_000, 999_999_999)
    attendant_phone = rng.randint(300_000_0000, 399_999_9999)
    attendant_email = (
        f"{email_part(attendant_first_name)[0]}{email_part(attendant_last_name)[0]}.acu"
        "@example.com"
    )

    # Remove the original four records while preserving all formatting and validations.
    for row in sheet.iter_rows(min_row=FIRST_DATA_ROW, max_row=sheet.max_row, min_col=1, max_col=18):
        for cell in row:
            cell.value = None

    generated_documents: set[int] = set()
    generated_emails: set[str] = set()

    for index in range(RECORD_COUNT):
        row_number = FIRST_DATA_ROW + index
        first_name = rng.choice(FIRST_NAMES)
        second_name = rng.choice(SECOND_NAMES) if rng.random() < 0.70 else None
        first_last_name = rng.choice(LAST_NAMES)
        second_last_name = rng.choice(LAST_NAMES) if rng.random() < 0.70 else None
        document_number = beneficiary_documents[index]
        phone_number = beneficiary_phones[index] if rng.random() < 0.85 else None
        student_email = None
        if rng.random() < 0.90:
            student_email = (
                f"{email_part(first_name)[0]}{email_part(first_last_name)[0]}{index + 1:03d}"
                "@example.com"
            )

        generated_documents.add(document_number)
        if student_email:
            generated_emails.add(student_email)

        values = [
            rng.choice(beneficiary_document_types),
            document_number,
            first_name,
            second_name,
            first_last_name,
            second_last_name,
            rng.choice(supported_grades),
            rng.choice(["a", "b", "c", "d"]),
            rng.choice(canonical_population_types),
            student_email,
            phone_number,
            "NIT",
            attendant_document,
            attendant_first_name,
            attendant_last_name,
            attendant_email,
            attendant_phone,
            None,
        ]

        for column_number, value in enumerate(values, start=1):
            sheet.cell(row=row_number, column=column_number, value=value)

    # Keep the workbook's original presentation and import mechanics.
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = "A2:R2"
    workbook.save(OUTPUT)

    if len(generated_documents) != RECORD_COUNT:
        raise AssertionError("Beneficiary document numbers are not unique")
    if len(generated_emails) != sum(
        1 for row in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1) if sheet.cell(row, 10).value
    ):
        raise AssertionError("Beneficiary email addresses are not unique")

    print(f"Created: {OUTPUT.resolve()}")
    print(f"Records: {RECORD_COUNT} (rows {FIRST_DATA_ROW}:{LAST_DATA_ROW})")
    print(
        "Repeated attendant:",
        attendant_document,
        attendant_first_name,
        attendant_last_name,
        attendant_phone,
    )

    verify_workbook(
        attendant=(
            "NIT",
            attendant_document,
            attendant_first_name,
            attendant_last_name,
            attendant_email,
            attendant_phone,
            None,
        )
    )


def verify_workbook(attendant: tuple[object, ...]) -> None:
    source_workbook = load_workbook(SOURCE, data_only=False)
    output_workbook = load_workbook(OUTPUT, data_only=False)
    source_sheet = source_workbook[source_workbook.sheetnames[0]]
    output_sheet = output_workbook[output_workbook.sheetnames[0]]

    source_headers = [source_sheet.cell(2, column).value for column in range(1, 19)]
    output_headers = [output_sheet.cell(2, column).value for column in range(1, 19)]
    assert output_headers == source_headers, "Headers changed"
    assert [str(item) for item in output_sheet.merged_cells.ranges] == ["A1:K1", "L1:R1"]
    assert output_sheet.freeze_panes == "A3"
    assert output_sheet.auto_filter.ref == "A2:R2"

    document_types = set(validation_values(output_sheet, "A3"))
    grades = {int(value) for value in validation_values(output_sheet, "G3")}
    populations = set(validation_values(output_sheet, "I3"))
    assert not document_types.intersection(UNSUPPORTED_DOCUMENT_TYPES)
    assert not grades.intersection(UNSUPPORTED_GRADES)
    assert set(validation_values(source_sheet, "I3")) == populations
    assert validation_values(output_sheet, "R3") == validation_values(source_sheet, "R3")
    required_columns = [1, 2, 3, 5, 7, 8, 9, 12, 13, 14, 15, 17]
    beneficiary_ids: list[int] = []
    attendant_rows: list[tuple[object, ...]] = []

    for row_number in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
        row = [output_sheet.cell(row_number, column).value for column in range(1, 19)]
        assert all(row[column - 1] not in (None, "") for column in required_columns)
        assert row[0] in document_types
        assert row[6] in grades
        assert row[8] in populations
        assert isinstance(row[1], int)
        if row[9] is not None:
            assert re.fullmatch(r"[a-z0-9.]+@example\.com", row[9])
        if row[10] is not None:
            assert isinstance(row[10], int) and len(str(row[10])) == 10
        assert isinstance(row[12], int)
        assert isinstance(row[16], int) and len(str(row[16])) == 10
        beneficiary_ids.append(row[1])
        attendant_rows.append(tuple(row[11:18]))

    assert len(beneficiary_ids) == RECORD_COUNT
    assert len(set(beneficiary_ids)) == RECORD_COUNT
    assert set(attendant_rows) == {attendant}
    assert not any(
        output_sheet.cell(row, column).value is not None
        for row in range(LAST_DATA_ROW + 1, output_sheet.max_row + 1)
        for column in range(1, 19)
    )
    assert not any(
        isinstance(cell.value, str) and cell.value.startswith("#")
        for row in output_sheet.iter_rows()
        for cell in row
    )
    print(
        "Verification: 100 records, unique beneficiary documents, one repeated attendant, "
        "unsupported document types and grades excluded"
    )


if __name__ == "__main__":
    main()
