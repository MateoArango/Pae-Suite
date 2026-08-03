from __future__ import annotations

from pathlib import Path
from openpyxl import load_workbook


SOURCE = Path("fixtures/student-data-alright - repeated - attendant  - x4.xlsx")


def main() -> None:
    workbook = load_workbook(SOURCE, data_only=False)
    print(f"Workbook: {SOURCE}")
    print(f"Sheets: {workbook.sheetnames}")
    print(f"Defined names: {[name.name for name in workbook.defined_names.values()]}")

    for sheet in workbook.worksheets:
        print("\n" + "=" * 80)
        print(
            f"Sheet={sheet.title!r} state={sheet.sheet_state} "
            f"dimension={sheet.calculate_dimension()} max_row={sheet.max_row} max_col={sheet.max_column}"
        )
        print(f"Merged ranges: {[str(item) for item in sheet.merged_cells.ranges]}")
        print(f"Freeze panes: {sheet.freeze_panes}")
        print(f"Auto filter: {sheet.auto_filter.ref}")
        print(f"Tables: {list(sheet.tables.keys())}")
        print("Column widths:", {key: value.width for key, value in sheet.column_dimensions.items() if value.width})
        print("Row heights:", {key: value.height for key, value in sheet.row_dimensions.items() if value.height})

        for row in sheet.iter_rows():
            populated = []
            for cell in row:
                if cell.value is not None:
                    populated.append(
                        {
                            "coordinate": cell.coordinate,
                            "value": cell.value,
                            "data_type": cell.data_type,
                            "number_format": cell.number_format,
                            "style_id": cell.style_id,
                            "comment": cell.comment.text if cell.comment else None,
                        }
                    )
            if populated:
                print(populated)

        validations = list(sheet.data_validations.dataValidation)
        print(f"Data validations ({len(validations)}):")
        for validation in validations:
            print(
                {
                    "sqref": str(validation.sqref),
                    "type": validation.type,
                    "operator": validation.operator,
                    "formula1": validation.formula1,
                    "formula2": validation.formula2,
                    "allow_blank": validation.allow_blank,
                    "error": validation.error,
                    "prompt": validation.prompt,
                }
            )


if __name__ == "__main__":
    main()
