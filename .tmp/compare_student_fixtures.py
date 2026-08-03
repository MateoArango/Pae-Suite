from pathlib import Path

from openpyxl import load_workbook


for path in sorted(Path("fixtures").glob("*.xlsx")):
    if path.name.startswith("~$"):
        continue
    workbook = load_workbook(path, read_only=False, data_only=False)
    sheet = workbook[workbook.sheetnames[0]]
    print("\n" + "=" * 100)
    print(path.name, sheet.calculate_dimension(), "rows", sheet.max_row, "cols", sheet.max_column)
    for row_number in range(1, min(sheet.max_row, 12) + 1):
        values = [sheet.cell(row_number, column).value for column in range(1, 19)]
        if any(value is not None for value in values):
            print(row_number, values)
    print(
        "validations",
        [
            (str(item.sqref), item.type, item.formula1)
            for item in sheet.data_validations.dataValidation
        ],
    )
